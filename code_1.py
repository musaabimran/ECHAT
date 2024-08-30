import streamlit as st
from langchain_community.vectorstores import Chroma
from langchain_community.chat_models import ChatOllama
from langchain_community.embeddings import FastEmbedEmbeddings
from langchain.schema.output_parser import StrOutputParser
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema.runnable import RunnablePassthrough
from langchain.prompts import PromptTemplate
from langchain.vectorstores.utils import filter_complex_metadata
from langchain.docstore.document import Document
import openpyxl
import os

class ChatExcel:
    def __init__(self):
        self.vector_store = None
        self.retriever = None
        self.chain = None
        self.text_splitter = RecursiveCharacterTextSplitter(chunk_size=1024, chunk_overlap=100)
        self.prompt = PromptTemplate.from_template(
            """
            <s> [INST] You are an assistant for answering questions. Use the following context to answer the question.
            If you don't know the answer, just say you don't know. Use a maximum of three sentences and be concise in your response. [/INST] </s>
            [INST] Question: {question}
            Context: {context}
            Answer: [/INST]
            """
        )

    def load_data_from_excel(self, file_path):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        documents = []
        for row in sheet.iter_rows(values_only=True):
            text = " ".join([str(cell) for cell in row if cell is not None])
            documents.append(Document(page_content=text))
        return documents

    def ingest(self, excel_file_path: str):
        try:
            docs = self.load_data_from_excel(excel_file_path)
            chunks = self.text_splitter.split_documents(docs)
            chunks = filter_complex_metadata(chunks)

            self.vector_store = Chroma.from_documents(documents=chunks, embedding=FastEmbedEmbeddings())
            self.retriever = self.vector_store.as_retriever(
                search_type="similarity_score_threshold",
                search_kwargs={
                    "k": 3,
                    "score_threshold": 0.5,
                },
            )

            self.chain = ({"context": self.retriever, "question": RunnablePassthrough()}
                          | self.prompt
                          | ChatOllama(model="mistral")
                          | StrOutputParser())
        except Exception as e:
            st.error(f"Error during ingestion: {e}")

    def ask(self, query: str):
        if not self.chain:
            return "Please, add an Excel document first."
        try:
            return self.chain.invoke(query)
        except Exception as e:
            return f"Error during querying: {e}"

    def clear(self):
        self.vector_store = None
        self.retriever = None
        self.chain = None

st.title("Excel Chatbot")

chat_excel = ChatExcel()

# Create temp directory if it doesn't exist
if not os.path.exists("temp"):
    os.makedirs("temp")

uploaded_file = st.file_uploader("Upload an Excel file", type=["xlsx"])
if uploaded_file:
    file_path = os.path.join("temp", uploaded_file.name)
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    st.success("File uploaded successfully.")
    chat_excel.ingest(file_path)

query = st.text_input("Ask a question about the Excel file:")
if st.button("Submit"):
    if uploaded_file and query:
        response = chat_excel.ask(query)
        st.write(response)
    else:
        st.error("Please upload an Excel file and enter a query.")

if st.button("Clear"):
    chat_excel.clear()
    st.success("Chatbot cleared.")
